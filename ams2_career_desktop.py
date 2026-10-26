import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import json
import sqlite3
import os
import random
import openpyxl
import ctypes
import sys

try:
    # Forces Windows taskbar to treat this as its own unique application
    ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("AMS2CareerManager.Desktop.App")
except Exception:
    pass

# SECURE APPDATA PERSISTENCE SETUP
APPDATA_DIR = os.path.join(os.environ.get('APPDATA', os.path.expanduser('~')), 'AMS2CareerManager')
os.makedirs(APPDATA_DIR, exist_ok=True)
DB_FILE = os.path.join(APPDATA_DIR, "ams2_careers.db")

# BLUEPRINTS DIRECTORY SETUP
BLUEPRINTS_DIR = os.path.join(APPDATA_DIR, "blueprints")
os.makedirs(BLUEPRINTS_DIR, exist_ok=True)

def init_db():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS careers (
            id TEXT PRIMARY KEY,
            name TEXT,
            driver_name TEXT,
            season INTEGER,
            current_contract TEXT,
            races TEXT,
            calendar TEXT,
            history TEXT,
            blueprint_name TEXT DEFAULT ''
        )
    """)
    c.execute("PRAGMA table_info(careers)")
    columns = [col[1] for col in c.fetchall()]
    if "history" not in columns:
        try:
            c.execute("ALTER TABLE careers ADD COLUMN history TEXT DEFAULT '[]'")
        except Exception:
            pass
    if "blueprint_name" not in columns:
        try:
            c.execute("ALTER TABLE careers ADD COLUMN blueprint_name TEXT DEFAULT ''")
        except Exception:
            pass
    conn.commit()
    conn.close()

init_db()

class BlueprintEngine:
    def __init__(self):
        self.championships = []
        self.events = []
        self.cars = []
        self.is_loaded = False
        self.active_blueprint_filename = ""

    def load_from_excel(self, file_path):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        # 1. Load Championships & Meta Data
        if "championships" in wb.sheetnames:
            sheet = wb["championships"]
            headers = [str(cell.value).strip().lower() for cell in sheet[1] if cell.value is not None]
            name_idx = headers.index("name") if "name" in headers else 0
            cat_idx = headers.index("categories") if "categories" in headers else 1
            prest_idx = headers.index("prestige") if "prestige" in headers else 2
            events_idx = headers.index("events_count") if "events_count" in headers else 8
            
            start_idx = next((i for i, h in enumerate(headers) if "start" in h), None)
            field_idx = next((i for i, h in enumerate(headers) if "field" in h), None)
            parc_idx = next((i for i, h in enumerate(headers) if "parc" in h or "ferme" in h), None)
            tags_idx = next((i for i, h in enumerate(headers) if "tag" in h), None)
            
            self.championships = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[name_idx]: continue
                categories_raw = str(row[cat_idx]) if len(row) > cat_idx and row[cat_idx] else ""
                categories = [c.strip() for c in categories_raw.split(",") if c.strip()]
                
                start_type = str(row[start_idx]).strip() if start_idx is not None and len(row) > start_idx and row[start_idx] else "Rolling"
                field_type_val = str(row[field_idx]).strip() if field_idx is not None and len(row) > field_idx and row[field_idx] else ""
                field_type = "Identical" if field_type_val and field_type_val.lower() != "none" else "Non-Identical"
                
                parc_raw = str(row[parc_idx]).strip().lower() if parc_idx is not None and len(row) > parc_idx and row[parc_idx] is not None else "false"
                parc_ferme = "True" if parc_raw in ["true", "1", "yes"] else "False"

                tags_raw = str(row[tags_idx]).strip() if tags_idx is not None and len(row) > tags_idx and row[tags_idx] else ""
                tags = [t.strip() for t in tags_raw.split(",") if t.strip()]

                self.championships.append({
                    "name": str(row[name_idx]).strip(),
                    "categories": categories,
                    "tags": tags,
                    "prestige": int(row[prest_idx] or 50),
                    "events_count": int(row[events_idx] or 8),
                    "start_type": start_type,
                    "field_type": field_type,
                    "parc_ferme": parc_ferme
                })

        # 2. Load Events (Tracks & Variants)
        if "events" in wb.sheetnames:
            sheet = wb["events"]
            self.events = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[0]: continue
                champ_name = str(row[0]).strip()
                track_name = str(row[1]).strip()
                variant = str(row[2]).strip() if len(row) > 2 and row[2] else "Grand Prix"
                self.events.append({
                    "championship": champ_name,
                    "track_name": track_name,
                    "variant": variant
                })

        # 3. Load Cars & Teams
        if "cars" in wb.sheetnames:
            sheet = wb["cars"]
            headers = [str(cell.value).strip().lower() for cell in sheet[1] if cell.value is not None]
            
            t_idx = next((i for i, h in enumerate(headers) if "team" in h), 0)
            c_idx = next((i for i, h in enumerate(headers) if "cat" in h), 1)
            m_idx = next((i for i, h in enumerate(headers) if "model" in h or "car" in h), 2 if len(headers) > 2 else 0)

            self.cars = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[0]: continue
                team_name = str(row[t_idx]).strip() if len(row) > t_idx and row[t_idx] else "Factory Squad"
                category = str(row[c_idx]).strip() if len(row) > c_idx and row[c_idx] else "General"
                car_model = str(row[m_idx]).strip() if len(row) > m_idx and row[m_idx] else "Spec Race Car"
                
                self.cars.append({
                    "team_name": team_name,
                    "category": category,
                    "car_model": car_model
                })

        self.is_loaded = True
        self.active_blueprint_filename = os.path.basename(file_path)
        return len(self.championships)

    def generate_calendar_for_champ(self, champ_name):
        champ = next((c for c in self.championships if c['name'].lower() == champ_name.lower()), None)
        target_count = champ['events_count'] if champ else 6
        
        matching_events = [e for e in self.events if e['championship'].lower() == champ_name.lower()]
        
        calendar = []
        if matching_events:
            for i in range(target_count):
                ev = matching_events[i % len(matching_events)]
                calendar.append(f"{ev['track_name']} ({ev['variant']})")
        else:
            calendar = [f"Standard Circuit R{i+1}" for i in range(target_count)]
            
        return calendar

engine = BlueprintEngine()

class AMS2CareerApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("AMS2 Career Manager")
        self.geometry("1260x780")
        self.configure(bg="#090d16")
        
        # BULLETPROOF RUNTIME ICON LOADER & WIN32 TASKBAR FORCE MAPPER
        self.set_window_icon()
        
        self.style = ttk.Style(self)
        self.style.theme_use("clam")
        
        self.style.configure("Treeview", 
                             background="#1e293b", 
                             foreground="#f8fafc", 
                             fieldbackground="#1e293b",
                             rowheight=28,
                             bordercolor="#334155",
                             lightcolor="#334155",
                             darkcolor="#334155",
                             borderwidth=1)
        self.style.map("Treeview", 
                       background=[("selected", "#2563eb")], 
                       foreground=[("selected", "#ffffff")])
        
        self.style.configure("Treeview.Heading", 
                             background="#0f172a", 
                             foreground="#94a3b8", 
                             font=("Helvetica", 9, "bold"),
                             relief=tk.SOLID,
                             bordercolor="#334155",
                             darkcolor="#334155",
                             lightcolor="#334155",
                             borderwidth=1)
        self.style.map("Treeview.Heading", background=[("active", "#1e293b")])
        
        self.style.configure("TCombobox", fieldbackground="#1e293b", background="#334155", foreground="#ffffff")
        self.style.configure("TSpinbox", fieldbackground="#1e293b", background="#334155", foreground="#ffffff")
        self.style.configure("TEntry", fieldbackground="#1e293b", foreground="#ffffff")

        self.active_profile = None
        self.create_widgets()
        self.load_profiles_list()

    def set_window_icon(self):
        try:
            if hasattr(sys, '_MEIPASS'):
                icon_path = os.path.join(sys._MEIPASS, "app_icon.ico")
            else:
                icon_path = os.path.abspath("app_icon.ico")
            
            if os.path.exists(icon_path):
                self.iconbitmap(icon_path)
                
                # Force-apply directly to Win32 HWND for taskbar certainty
                self.update_idletasks()
                hwnd = ctypes.windll.user32.GetParent(self.winfo_id())
                if not hwnd:
                    hwnd = self.winfo_id()
                
                IMAGE_ICON = 1
                LR_LOADFROMFILE = 0x00000010
                hicon_large = ctypes.windll.user32.LoadImageW(0, icon_path, IMAGE_ICON, 32, 32, LR_LOADFROMFILE)
                hicon_small = ctypes.windll.user32.LoadImageW(0, icon_path, IMAGE_ICON, 16, 16, LR_LOADFROMFILE)
                
                WM_SETICON = 0x0080
                ICON_SMALL = 0
                ICON_BIG = 1
                
                if hicon_large:
                    ctypes.windll.user32.SendMessageW(hwnd, WM_SETICON, ICON_BIG, hicon_large)
                if hicon_small:
                    ctypes.windll.user32.SendMessageW(hwnd, WM_SETICON, ICON_SMALL, hicon_small)
        except Exception:
            pass

    def create_widgets(self):
        header = tk.Frame(self, bg="#030712", height=60)
        header.pack(fill=tk.X, side=tk.TOP)
        
        tk.Label(header, text="AMS2 Career Manager", font=("Helvetica", 15, "bold"), fg="#f8fafc", bg="#030712").pack(side=tk.LEFT, padx=20, pady=12)
        
        btn_frame = tk.Frame(header, bg="#030712")
        btn_frame.pack(side=tk.RIGHT, padx=20)
        
        tk.Button(btn_frame, text="Career History", command=self.show_history_screen, bg="#334155", fg="white", relief=tk.FLAT, font=("Helvetica", 9, "bold"), padx=10, pady=5).pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="Switch Career", command=self.show_select_screen, bg="#1e293b", fg="white", relief=tk.FLAT, font=("Helvetica", 9, "bold"), padx=10, pady=5).pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="Import Blueprint (.xlsx)", command=self.import_blueprint_file, bg="#059669", fg="white", relief=tk.FLAT, font=("Helvetica", 9, "bold"), padx=10, pady=5).pack(side=tk.LEFT, padx=5)

        self.container = tk.Frame(self, bg="#090d16")
        self.container.pack(fill=tk.BOTH, expand=True)

        self.screens = {}
        self.build_select_screen()
        self.build_dashboard_screen()
        self.build_summary_screen()
        self.build_history_screen()

        self.show_select_screen()

    def show_screen(self, name):
        for s in self.screens.values(): s.pack_forget()
        self.screens[name].pack(fill=tk.BOTH, expand=True)

    # ------------------ SCREEN 1: CAREER SELECT ------------------
    def build_select_screen(self):
        frame = tk.Frame(self.container, bg="#090d16", padx=30, pady=25)
        self.screens['select'] = frame

        tk.Label(frame, text="Saved Careers", font=("Helvetica", 18, "bold"), bg="#090d16", fg="#f8fafc").pack(anchor="w")
        
        # BLUEPRINTS LIBRARY CONTROLS BAR
        bp_bar = tk.Frame(frame, bg="#1e293b", padx=10, pady=8)
        bp_bar.pack(fill=tk.X, pady=(6, 12))

        tk.Label(bp_bar, text="Blueprint Library:", font=("Helvetica", 9, "bold"), bg="#1e293b", fg="#94a3b8").pack(side=tk.LEFT, padx=(0, 8))
        
        self.bp_combo_var = tk.StringVar()
        self.bp_combo = ttk.Combobox(bp_bar, textvariable=self.bp_combo_var, state="readonly", width=35)
        self.bp_combo.pack(side=tk.LEFT, padx=5)
        self.bp_combo.bind("<<ComboboxSelected>>", self.on_select_blueprint_from_library)

        tk.Button(bp_bar, text="Refresh Library", command=self.refresh_blueprint_library, bg="#334155", fg="white", font=("Helvetica", 8, "bold"), padx=8, relief=tk.FLAT).pack(side=tk.LEFT, padx=5)

        # FULL UNCLIPPED FILENAME DISPLAY LABEL (TOOLTIP / LABEL)
        self.bp_file_lbl = tk.Label(bp_bar, text="Active File: None", font=("Helvetica", 9, "italic"), bg="#1e293b", fg="#38bdf8")
        self.bp_file_lbl.pack(side=tk.LEFT, padx=15)
        
        self.status_lbl = tk.Label(frame, text="Blueprint Status: Not Loaded", font=("Helvetica", 9, "bold"), bg="#090d16", fg="#ef4444")
        self.status_lbl.pack(anchor="w", pady=(2, 10))

        btn_box = tk.Frame(frame, bg="#090d16")
        btn_box.pack(fill=tk.X, pady=(0, 10))
        
        tk.Button(btn_box, text="+ New Career Profile", command=self.open_new_career_dialog, bg="#2563eb", fg="white", font=("Helvetica", 10, "bold"), padx=15, pady=8, relief=tk.FLAT).pack(side=tk.LEFT)

        list_frame = tk.Frame(frame, bg="#1e293b", highlightbackground="#334155", highlightthickness=1)
        list_frame.pack(fill=tk.BOTH, expand=True)

        self.tree = ttk.Treeview(list_frame, columns=("driver", "season", "team", "class", "blueprint"), show="headings", height=10)
        self.tree.heading("driver", text="Driver Name")
        self.tree.heading("season", text="Current Season")
        self.tree.heading("team", text="Current Team / Car")
        self.tree.heading("class", text="Championship Series")
        self.tree.heading("blueprint", text="Source Blueprint")
        
        self.tree.column("driver", width=140, stretch=False)
        self.tree.column("season", width=100, stretch=False)
        self.tree.column("team", width=240, stretch=False)
        self.tree.column("class", width=220, stretch=False)
        self.tree.column("blueprint", width=200, stretch=True)

        self.tree.pack(fill=tk.BOTH, expand=True)

        action_frame = tk.Frame(frame, bg="#090d16", pady=15)
        action_frame.pack(fill=tk.X)
        tk.Button(action_frame, text="Load Selected Career", command=self.load_selected_career, bg="#059669", fg="white", font=("Helvetica", 10, "bold"), padx=20, pady=8, relief=tk.FLAT).pack(side=tk.LEFT)

        self.refresh_blueprint_library()

    def import_blueprint_file(self):
        f_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        if f_path:
            try:
                import shutil
                filename = os.path.basename(f_path)
                dest_path = os.path.join(BLUEPRINTS_DIR, filename)
                shutil.copy(f_path, dest_path)
                
                count = engine.load_from_excel(dest_path)
                self.refresh_blueprint_library()
                self.bp_combo.set(filename)
                self.update_blueprint_width_and_label(filename)
                self.status_lbl.config(text=f"Blueprint Status: Loaded '{filename}' ({count} Championships)", fg="#10b981")
                messagebox.showinfo("Success", f"Imported and loaded '{filename}' successfully!")
            except Exception as e:
                messagebox.showerror("Error Importing Excel", f"Could not import blueprint workbook:\n{e}")

    def refresh_blueprint_library(self):
        if not os.path.exists(BLUEPRINTS_DIR):
            os.makedirs(BLUEPRINTS_DIR, exist_ok=True)
        files = [f for f in os.listdir(BLUEPRINTS_DIR) if f.endswith('.xlsx')]
        self.bp_combo['values'] = files
        if files:
            if engine.active_blueprint_filename in files:
                target_file = engine.active_blueprint_filename
            else:
                target_file = files[0]
            self.bp_combo.set(target_file)
            self.load_blueprint_by_filename(target_file)
        else:
            self.bp_combo.set('')
            self.bp_file_lbl.config(text="Active File: None")
            self.status_lbl.config(text="Blueprint Status: No Blueprints Found in Library (Import above)", fg="#ef4444")

    def update_blueprint_width_and_label(self, filename):
        dynamic_width = max(35, min(len(filename) + 5, 55))
        self.bp_combo.config(width=dynamic_width)
        self.bp_file_lbl.config(text=f"Active File: {filename}")

    def on_select_blueprint_from_library(self, event):
        filename = self.bp_combo.get()
        if filename:
            self.load_blueprint_by_filename(filename)

    def load_blueprint_by_filename(self, filename):
        f_path = os.path.join(BLUEPRINTS_DIR, filename)
        if os.path.exists(f_path):
            try:
                count = engine.load_from_excel(f_path)
                self.update_blueprint_width_and_label(filename)
                self.status_lbl.config(text=f"Blueprint Status: Active '{filename}' ({count} Championships)", fg="#10b981")
            except Exception as e:
                messagebox.showerror("Error", f"Could not load blueprint {filename}:\n{e}")

    def load_profiles_list(self):
        conn = sqlite3.connect(DB_FILE)
        c = conn.cursor()
        c.execute("SELECT id, name, driver_name, season, current_contract, blueprint_name FROM careers")
        rows = []
        for r in c.fetchall():
            contract = json.loads(r[4])
            bp_used = r[5] if len(r) > 5 and r[5] else "Default"
            rows.append((r[2], f"Season {r[3]}", f"{contract.get('teamName', 'N/A')} ({contract.get('carModel', 'N/A')})", contract.get('champName', 'N/A'), bp_used))
        conn.close()
        
        for i in self.tree.get_children(): self.tree.delete(i)
        conn = sqlite3.connect(DB_FILE)
        c = conn.cursor()
        c.execute("SELECT id FROM careers")
        ids = [row[0] for row in c.fetchall()]
        conn.close()

        for idx, row_vals in enumerate(rows):
            tag = "evenrow" if idx % 2 == 0 else "oddrow"
            self.tree.insert("", tk.END, iid=ids[idx], values=row_vals, tags=(tag,))
        self.tree.tag_configure("evenrow", background="#1e293b")
        self.tree.tag_configure("oddrow", background="#162032")

    def show_select_screen(self):
        self.load_profiles_list()
        self.show_screen('select')

    # ------------------ SCREEN 2: DASHBOARD ------------------
    def build_dashboard_screen(self):
        frame = tk.Frame(self.container, bg="#090d16", padx=25, pady=20)
        self.screens['dashboard'] = frame

        self.dash_header = tk.Label(frame, text="", font=("Helvetica", 14, "bold"), bg="#090d16", fg="#f8fafc", anchor="w")
        self.dash_header.pack(fill=tk.X, pady=(0, 6))

        self.meta_lbl = tk.Label(frame, text="", font=("Helvetica", 9, "bold"), bg="#1e293b", fg="#94a3b8", anchor="w", padx=10, pady=4)
        self.meta_lbl.pack(fill=tk.X, pady=(0, 10))

        content = tk.Frame(frame, bg="#090d16")
        content.pack(fill=tk.BOTH, expand=True)

        left_col = tk.Frame(content, bg="#1e293b", width=340, highlightbackground="#334155", highlightthickness=1, padx=15, pady=15)
        left_col.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 15))
        left_col.pack_propagate(False)

        tk.Label(left_col, text="Season Rank Status", font=("Helvetica", 10, "bold"), bg="#1e293b", fg="#f8fafc").pack(anchor="w")
        
        # LOCKED COLORS: GREEN, YELLOW, RED INDICATORS
        self.rank_box = tk.Frame(left_col, bg="#059669", padx=10, pady=10)
        self.rank_box.pack(fill=tk.X, pady=(5, 15))
        
        self.rank_title_lbl = tk.Label(self.rank_box, text="PROMOTION", font=("Helvetica", 11, "bold"), bg="#059669", fg="white")
        self.rank_title_lbl.pack()
        self.rank_sub_lbl = tk.Label(self.rank_box, text="Trajectory: 85.0%", font=("Helvetica", 9), bg="#059669", fg="white")
        self.rank_sub_lbl.pack(pady=2)

        tk.Label(left_col, text="Season Event Calendar", font=("Helvetica", 10, "bold"), bg="#1e293b", fg="#f8fafc").pack(anchor="w")
        
        cal_frame = tk.Frame(left_col, bg="#1e293b")
        cal_frame.pack(fill=tk.BOTH, expand=True, pady=8)
        
        self.cal_listbox = tk.Listbox(cal_frame, font=("Helvetica", 8), bg="#090d16", fg="#f8fafc", selectbackground="#2563eb", selectforeground="#ffffff", borderwidth=0, highlightthickness=0)
        self.cal_listbox.pack(fill=tk.BOTH, expand=True)

        tk.Button(left_col, text="End Season & Evaluate Offers", command=self.evaluate_season, bg="#0f172a", fg="white", font=("Helvetica", 9, "bold"), pady=8, relief=tk.FLAT).pack(fill=tk.X, side=tk.BOTTOM)

        right_col = tk.Frame(content, bg="#090d16")
        right_col.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        form_box = tk.LabelFrame(right_col, text=" Log Race Result ", font=("Helvetica", 10, "bold"), bg="#1e293b", fg="#f8fafc", padx=12, pady=12)
        form_box.pack(fill=tk.X, pady=(0, 12))

        f_grid = tk.Frame(form_box, bg="#1e293b")
        f_grid.pack(fill=tk.X)

        tk.Label(f_grid, text="Round / Track Variant", bg="#1e293b", fg="#94a3b8", font=("Helvetica", 8, "bold")).grid(row=0, column=0, padx=4, sticky="w")
        self.round_combo = ttk.Combobox(f_grid, width=24, state="readonly")
        self.round_combo.grid(row=1, column=0, padx=4, pady=(0, 5))

        tk.Label(f_grid, text="Quali", bg="#1e293b", fg="#94a3b8", font=("Helvetica", 8, "bold")).grid(row=0, column=1, padx=4, sticky="w")
        self.ent_q = ttk.Spinbox(f_grid, from_=1, to=100, width=5)
        self.ent_q.grid(row=1, column=1, padx=4, pady=(0, 5))
        self.ent_q.set(1)

        tk.Label(f_grid, text="Finish", bg="#1e293b", fg="#94a3b8", font=("Helvetica", 8, "bold")).grid(row=0, column=2, padx=4, sticky="w")
        self.ent_f = ttk.Spinbox(f_grid, from_=1, to=100, width=5)
        self.ent_f.grid(row=1, column=2, padx=4, pady=(0, 5))
        self.ent_f.set(1)

        tk.Label(f_grid, text="Champ Pos", bg="#1e293b", fg="#94a3b8", font=("Helvetica", 8, "bold")).grid(row=0, column=3, padx=4, sticky="w")
        self.ent_c = ttk.Spinbox(f_grid, from_=1, to=100, width=5)
        self.ent_c.grid(row=1, column=3, padx=4, pady=(0, 5))
        self.ent_c.set(1)

        tk.Label(f_grid, text="Grid Size", bg="#1e293b", fg="#94a3b8", font=("Helvetica", 8, "bold")).grid(row=0, column=4, padx=4, sticky="w")
        self.ent_grid = ttk.Spinbox(f_grid, from_=2, to=60, width=5)
        self.ent_grid.grid(row=1, column=4, padx=4, pady=(0, 5))
        self.ent_grid.set(20)

        tk.Button(form_box, text="+ Record Result", command=self.log_race, bg="#2563eb", fg="white", font=("Helvetica", 9, "bold"), padx=12, pady=4, relief=tk.FLAT).pack(anchor="e", pady=(4, 0))

        log_box = tk.Frame(right_col, bg="#1e293b", highlightbackground="#334155", highlightthickness=1)
        log_box.pack(fill=tk.BOTH, expand=True)

        self.race_tree = ttk.Treeview(log_box, columns=("event", "q", "f", "c", "grid"), show="headings", height=7)
        self.race_tree.heading("event", text="Track & Layout Variant")
        self.race_tree.heading("q", text="Quali")
        self.race_tree.heading("f", text="Finish")
        self.race_tree.heading("c", text="Champ Pos")
        self.race_tree.heading("grid", text="Grid Size")
        self.race_tree.pack(fill=tk.BOTH, expand=True, side=tk.TOP)

        edit_bar = tk.Frame(right_col, bg="#090d16", pady=8)
        edit_bar.pack(fill=tk.X)
        tk.Button(edit_bar, text="Edit Selected Result", command=self.edit_selected_race, bg="#d97706", fg="white", font=("Helvetica", 9, "bold"), padx=10, relief=tk.FLAT).pack(side=tk.LEFT, padx=(0, 10))
        tk.Button(edit_bar, text="Delete Selected Result", command=self.delete_selected_race, bg="#dc2626", fg="white", font=("Helvetica", 9, "bold"), padx=10, relief=tk.FLAT).pack(side=tk.LEFT)

    def calculate_ti(self):
        races = self.active_profile['races']
        if not races: return 0.5
        
        total_score = 0.0
        for r in races:
            g_size = max(2, int(r.get('grid', 20)))
            f_pos = max(1, min(r['f'], g_size))
            c_pos = max(1, min(r['c'], g_size))
            
            score_champ = max(0.0, 1.0 - ((c_pos - 1) / float(g_size - 1 if g_size > 1 else 1)))
            score_race = max(0.0, 1.0 - ((f_pos - 1) / float(g_size - 1 if g_size > 1 else 1)))
            
            total_score += (0.70 * score_champ) + (0.30 * score_race)
            
        return max(0.0, min(1.0, total_score / len(races)))

    def refresh_dashboard(self):
        p = self.active_profile
        contract = p['currentContract']
        champ_name = contract.get('champName')
        
        champ_obj = next((c for c in engine.championships if c['name'].lower() == champ_name.lower()), {})
        start_type = champ_obj.get('start_type', 'Rolling')
        field_type = champ_obj.get('field_type', 'Non-Identical')
        parc_ferme = champ_obj.get('parc_ferme', 'False')

        self.dash_header.config(text=f"{p['name']} | Season {p['season']} | {champ_name} ({contract.get('teamName')} - {contract.get('carModel')})")
        self.meta_lbl.config(text=f"Series Specs | Start: {start_type}  |  Field: {field_type}  |  Parc Fermé: {parc_ferme}  |  Prestige: {contract.get('prestige', 'N/A')}")

        self.cal_listbox.delete(0, tk.END)
        cal = p.get('calendar', [])
        for idx, ev in enumerate(cal):
            self.cal_listbox.insert(tk.END, f"R{idx+1}: {ev}")

        self.round_combo['values'] = [f"R{i+1}: {tr}" for i, tr in enumerate(cal)]
        if cal: self.round_combo.set(f"R1: {cal[0]}")

        race_rows = []
        for r in p['races']:
            race_rows.append((r['event'], f"P{r['q']}", f"P{r['f']}", f"P{r['c']}", r.get('grid', 20)))
        
        for i in self.race_tree.get_children(): self.race_tree.delete(i)
        for idx, row_vals in enumerate(race_rows):
            tag = "evenrow" if idx % 2 == 0 else "oddrow"
            self.race_tree.insert("", tk.END, iid=str(idx), values=row_vals, tags=(tag,))
        self.race_tree.tag_configure("evenrow", background="#1e293b")
        self.race_tree.tag_configure("oddrow", background="#162032")

        ti = self.calculate_ti()
        if ti >= 0.75:
            self.rank_box.config(bg="#059669")
            self.rank_title_lbl.config(text="PROMOTION", bg="#059669")
            self.rank_sub_lbl.config(text=f"Trajectory: {ti*100:.1f}%", bg="#059669")
        elif ti >= 0.40:
            self.rank_box.config(bg="#d97706")
            self.rank_title_lbl.config(text="RETAINED", bg="#d97706")
            self.rank_sub_lbl.config(text=f"Trajectory: {ti*100:.1f}%", bg="#d97706")
        else:
            self.rank_box.config(bg="#dc2626")
            self.rank_title_lbl.config(text="DEMOTION", bg="#dc2626")
            self.rank_sub_lbl.config(text=f"Trajectory: {ti*100:.1f}%", bg="#dc2626")

    def log_race(self):
        ev = self.round_combo.get()
        if not ev: return
        round_name_only = ev.split(":")[0]
        existing_rounds = [r['event'].split(":")[0] for r in self.active_profile['races']]
        if round_name_only in existing_rounds:
            messagebox.showwarning("Round Logged", f"{round_name_only} already logged.")
            return

        q = int(self.ent_q.get())
        f = int(self.ent_f.get())
        c = int(self.ent_c.get())
        grid = int(self.ent_grid.get())
        
        self.active_profile['races'].append({"event": ev, "q": q, "f": f, "c": c, "grid": grid})
        self.save_active_profile()
        self.refresh_dashboard()

    def delete_selected_race(self):
        sel = self.race_tree.selection()
        if not sel: return
        del self.active_profile['races'][int(sel[0])]
        self.save_active_profile()
        self.refresh_dashboard()

    def edit_selected_race(self):
        sel = self.race_tree.selection()
        if not sel: return
        idx = int(sel[0])
        r = self.active_profile['races'][idx]

        dlg = tk.Toplevel(self)
        dlg.title("Edit Race Result")
        dlg.geometry("300x290")
        dlg.configure(bg="#090d16")

        tk.Label(dlg, text=f"Editing: {r['event']}", bg="#090d16", fg="#f8fafc").pack(pady=10)
        tk.Label(dlg, text="Quali Position:", bg="#090d16", fg="#94a3b8").pack(anchor="w", padx=20)
        eq = ttk.Spinbox(dlg, from_=1, to=100, width=10); eq.set(r['q']); eq.pack(padx=20, pady=2)
        tk.Label(dlg, text="Finish Position:", bg="#090d16", fg="#94a3b8").pack(anchor="w", padx=20)
        ef = ttk.Spinbox(dlg, from_=1, to=100, width=10); ef.set(r['f']); ef.pack(padx=20, pady=2)
        tk.Label(dlg, text="Champ Position:", bg="#090d16", fg="#94a3b8").pack(anchor="w", padx=20)
        ec = ttk.Spinbox(dlg, from_=1, to=100, width=10); ec.set(r['c']); ec.pack(padx=20, pady=2)
        tk.Label(dlg, text="Grid Size:", bg="#090d16", fg="#94a3b8").pack(anchor="w", padx=20)
        eg = ttk.Spinbox(dlg, from_=2, to=60, width=10); eg.set(r.get('grid', 20)); eg.pack(padx=20, pady=2)

        def save_edit():
            r['q'] = int(eq.get())
            r['f'] = int(ef.get())
            r['c'] = int(ec.get())
            r['grid'] = int(eg.get())
            self.save_active_profile()
            self.refresh_dashboard()
            dlg.destroy()

        tk.Button(dlg, text="Save Changes", command=save_edit, bg="#2563eb", fg="white", font=("Helvetica", 9, "bold"), relief=tk.FLAT).pack(pady=15)

    def save_active_profile(self):
        p = self.active_profile
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        cursor.execute("UPDATE careers SET season = ?, current_contract = ?, races = ?, calendar = ?, history = ?, blueprint_name = ? WHERE id = ?",
                       (p['season'], json.dumps(p['currentContract']), json.dumps(p['races']), json.dumps(p['calendar']), json.dumps(p.get('history', [])), p.get('blueprint_name', ''), p['id']))
        conn.commit()
        conn.close()

    # ------------------ SCREEN 3: SUMMARY & OFFERS ------------------
    def build_summary_screen(self):
        frame = tk.Frame(self.container, bg="#090d16", padx=25, pady=20)
        self.screens['summary'] = frame

        self.sum_title = tk.Label(frame, text="", font=("Helvetica", 14, "bold"), bg="#090d16", fg="#f8fafc")
        self.sum_title.pack(anchor="w")

        tk.Label(frame, text="Contract Market Offers (Full Continuum Range, Car Classes & Variable Seats)", font=("Helvetica", 10, "bold"), bg="#090d16", fg="#94a3b8").pack(anchor="w", pady=(6, 10))

        offer_container = tk.Frame(frame, bg="#1e293b", highlightbackground="#334155", highlightthickness=1)
        offer_container.pack(fill=tk.BOTH, expand=True)

        self.offer_tree = ttk.Treeview(offer_container, columns=("tier_type", "series", "car_class", "team", "car", "tags", "prestige"), show="headings", height=12)
        self.offer_tree.heading("tier_type", text="Offer Type")
        self.offer_tree.heading("series", text="Championship Series")
        self.offer_tree.heading("car_class", text="Car Class")
        self.offer_tree.heading("team", text="Team Name")
        self.offer_tree.heading("car", text="Car Model")
        self.offer_tree.heading("tags", text="Tags")
        self.offer_tree.heading("prestige", text="Series Prestige")
        
        self.offer_tree.column("tier_type", width=150, stretch=False)
        self.offer_tree.column("series", width=200, stretch=False)
        self.offer_tree.column("car_class", width=100, stretch=False)
        self.offer_tree.column("team", width=180, stretch=False)
        self.offer_tree.column("car", width=180, stretch=False)
        self.offer_tree.column("tags", width=180, stretch=True)
        self.offer_tree.column("prestige", width=100, stretch=False)

        h_scroll = ttk.Scrollbar(offer_container, orient="horizontal", command=self.offer_tree.xview)
        v_scroll = ttk.Scrollbar(offer_container, orient="vertical", command=self.offer_tree.yview)
        self.offer_tree.configure(xscrollcommand=h_scroll.set, yscrollcommand=v_scroll.set)

        self.offer_tree.grid(row=0, column=0, sticky="nsew")
        v_scroll.grid(row=0, column=1, sticky="ns")
        h_scroll.grid(row=1, column=0, sticky="ew")

        offer_container.grid_rowconfigure(0, weight=1)
        offer_container.grid_columnconfigure(0, weight=1)

        tk.Button(frame, text="Sign Selected Contract Offer", command=self.sign_contract, bg="#2563eb", fg="white", font=("Helvetica", 10, "bold"), pady=8, relief=tk.FLAT).pack(fill=tk.X, pady=15)

    def evaluate_season(self):
        self.sum_title.config(text=f"Season {self.active_profile['season']} Market Evaluation")
        for i in self.offer_tree.get_children(): self.offer_tree.delete(i)

        current_champ_name = self.active_profile['currentContract'].get('champName', '')
        current_champ = next((c for c in engine.championships if c['name'].lower() == current_champ_name.lower()), {"prestige": 5, "name": current_champ_name, "categories": [], "tags": []})
        current_prestige = current_champ['prestige']
        
        ti = self.calculate_ti()
        offers_pool = []

        sorted_champs = sorted(engine.championships, key=lambda x: x['prestige'])

        if ti >= 0.75:
            offers_pool.append(("Retention (Current Tier)", current_champ))

            max_prestige_jump = int(5 + ((ti - 0.75) / 0.25) * 10)
            target_max_prestige = current_prestige + max_prestige_jump

            promotion_candidates = [c for c in sorted_champs if current_prestige < c['prestige'] <= target_max_prestige]
            
            for c in promotion_candidates:
                offers_pool.append(("Promotion (Higher Tier)", c))
            
            lateral_champs = [c for c in sorted_champs if c['prestige'] == current_prestige and c['name'].lower() != current_champ_name.lower()]
            for c in lateral_champs[:2]:
                offers_pool.append(("Lateral (Same Tier)", c))

            fallback_champs = [c for c in sorted_champs if c['prestige'] < current_prestige]
            if fallback_champs:
                offers_pool.append(("Fallback (Lower Tier)", max(fallback_champs, key=lambda x: x['prestige'])))

        elif ti >= 0.40:
            retention_count = int(1 + (ti * 2))
            for _ in range(max(1, retention_count)):
                offers_pool.append(("Retention (Current Tier)", current_champ))

            lateral_champs = [c for c in sorted_champs if c['prestige'] == current_prestige and c['name'].lower() != current_champ_name.lower()]
            for c in lateral_champs[:2]:
                offers_pool.append(("Lateral (Same Tier)", c))

            fallback_champs = [c for c in sorted_champs if c['prestige'] < current_prestige]
            if fallback_champs:
                for c in sorted(fallback_champs, key=lambda x: x['prestige'], reverse=True)[:2]:
                    offers_pool.append(("Fallback (Lower Tier)", c))

        else:
            drop_depth = int(5 + ((1.0 - ti) * 20))
            fallback_champs = [c for c in sorted_champs if c['prestige'] < current_prestige and c['prestige'] >= current_prestige - drop_depth]
            if not fallback_champs:
                fallback_champs = [c for c in sorted_champs if c['prestige'] < current_prestige]

            num_demotions = int(2 + ((1.0 - ti) * 3))
            if fallback_champs:
                selected_demotions = random.sample(fallback_champs, min(num_demotions, len(fallback_champs)))
                for c in selected_demotions:
                    offers_pool.append(("Demotion / Fallback Offer", c))
            else:
                offers_pool.append(("Fallback Offer", current_champ))

        self.current_offers = []
        
        def get_strictly_matched_cars(ch):
            champ_cats = [cat.lower().strip() for cat in ch.get('categories', [])]
            champ_name_lower = ch['name'].lower()
            
            matched = []
            for car in engine.cars:
                car_cat = car['category'].lower().strip()
                if any(cc == car_cat or cc in car_cat or car_cat in cc for cc in champ_cats):
                    matched.append(car)
            
            if not matched:
                for car in engine.cars:
                    car_cat = car['category'].lower().strip()
                    if any(w in car_cat for w in champ_name_lower.split() if len(w) > 3):
                        matched.append(car)
                        
            return matched if matched else [{"team_name": "Factory Squad", "category": "Spec", "car_model": "Spec Race Car"}]

        offer_rows = []
        for type_label, ch in offers_pool:
            available_cars = get_strictly_matched_cars(ch)
            
            roll = random.random()
            if ti >= 0.85:
                base_seats = 4 if roll > 0.3 else 5
            elif ti >= 0.50:
                base_seats = 3 if roll > 0.4 else (4 if roll > 0.1 else 2)
            else:
                base_seats = 2 if roll > 0.4 else 3
                
            num_team_seats = min(base_seats, len(available_cars))
            num_team_seats = max(1, num_team_seats)
            
            chosen_cars = random.sample(available_cars, num_team_seats)
            tags_str = ", ".join(ch['tags']) if ch['tags'] else "Standard"
            for car_info in chosen_cars:
                car_class = car_info.get('category', 'General')
                offer = {
                    "champName": ch['name'], 
                    "teamName": car_info['team_name'], 
                    "carModel": car_info['car_model'], 
                    "carClass": car_class,
                    "prestige": ch['prestige']
                }
                self.current_offers.append(offer)
                offer_rows.append((type_label, ch['name'], car_class, car_info['team_name'], car_info['car_model'], tags_str, f"Prestige {ch['prestige']}"))

        for idx, row_vals in enumerate(offer_rows):
            tag = "evenrow" if idx % 2 == 0 else "oddrow"
            self.offer_tree.insert("", tk.END, iid=str(idx), values=row_vals, tags=(tag,))
        self.offer_tree.tag_configure("evenrow", background="#1e293b")
        self.offer_tree.tag_configure("oddrow", background="#162032")

        self.show_screen('summary')

    def sign_contract(self):
        sel = self.offer_tree.selection()
        if not sel: return
        offer = self.current_offers[int(sel[0])]
        
        history_entry = {
            "season": self.active_profile['season'],
            "contract": self.active_profile['currentContract'],
            "ti": self.calculate_ti(),
            "races": self.active_profile['races']
        }
        if 'history' not in self.active_profile:
            self.active_profile['history'] = []
        self.active_profile['history'].append(history_entry)

        new_cal = engine.generate_calendar_for_champ(offer['champName'])

        self.active_profile['currentContract'] = offer
        self.active_profile['season'] += 1
        self.active_profile['races'] = []
        self.active_profile['calendar'] = new_cal
        self.save_active_profile()

        messagebox.showinfo("Signed!", f"Signed with {offer['teamName']} for {offer['champName']} ({offer.get('carClass', 'General')})!")
        self.refresh_dashboard()
        self.show_screen('dashboard')

    # ------------------ SCREEN 4: HISTORY VIEWER ------------------
    def build_history_screen(self):
        frame = tk.Frame(self.container, bg="#090d16", padx=25, pady=20)
        self.screens['history'] = frame

        tk.Label(frame, text="Career Profile History & Archive", font=("Helvetica", 14, "bold"), bg="#090d16", fg="#f8fafc").pack(anchor="w", pady=(0, 10))

        content = tk.Frame(frame, bg="#090d16")
        content.pack(fill=tk.BOTH, expand=True)

        left_col = tk.Frame(content, bg="#1e293b", width=340, highlightbackground="#334155", highlightthickness=1, padx=12, pady=12)
        left_col.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 15))
        left_col.pack_propagate(False)

        tk.Label(left_col, text="Completed Seasons", font=("Helvetica", 10, "bold"), bg="#1e293b", fg="#f8fafc").pack(anchor="w")

        self.hist_season_tree = ttk.Treeview(left_col, columns=("season", "series", "ti"), show="headings", height=12)
        self.hist_season_tree.heading("season", text="Season")
        self.hist_season_tree.heading("series", text="Championship")
        self.hist_season_tree.heading("ti", text="Score")
        
        self.hist_season_tree.column("season", width=70, stretch=False)
        self.hist_season_tree.column("series", width=160, stretch=False)
        self.hist_season_tree.column("ti", width=70, stretch=False)
        
        self.hist_season_tree.pack(fill=tk.BOTH, expand=True, pady=8)
        self.hist_season_tree.bind("<<TreeviewSelect>>", self.on_select_history_season)

        right_container = tk.Frame(content, bg="#1e293b", highlightbackground="#334155", highlightthickness=1)
        right_container.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        tk.Label(right_container, text="Season Race Log Breakdown", font=("Helvetica", 10, "bold"), bg="#1e293b", fg="#f8fafc", padx=12, pady=10).pack(anchor="w")
        
        self.hist_race_tree = ttk.Treeview(right_container, columns=("event", "q", "f", "c", "grid"), show="headings", height=10)
        self.hist_race_tree.heading("event", text="Track & Layout Variant")
        self.hist_race_tree.heading("q", text="Quali")
        self.hist_race_tree.heading("f", text="Finish")
        self.hist_race_tree.heading("c", text="Champ Pos")
        self.hist_race_tree.heading("grid", text="Grid Size")

        h_scroll_h = ttk.Scrollbar(right_container, orient="horizontal", command=self.hist_race_tree.xview)
        v_scroll_h = ttk.Scrollbar(right_container, orient="vertical", command=self.hist_race_tree.yview)
        self.hist_race_tree.configure(xscrollcommand=h_scroll_h.set, yscrollcommand=v_scroll_h.set)

        self.hist_race_tree.pack(fill=tk.BOTH, expand=True, padx=12, pady=(0, 4))

    def show_history_screen(self):
        if not self.active_profile: return
        for i in self.hist_season_tree.get_children(): self.hist_season_tree.delete(i)
        for i in self.hist_race_tree.get_children(): self.hist_race_tree.delete(i)

        history = self.active_profile.get('history', [])
        season_rows = []
        for h in history:
            season_rows.append((f"Season {h['season']}", h['contract'].get('champName'), f"{h['ti']*100:.1f}%"))
            
        for idx, row_vals in enumerate(season_rows):
            tag = "evenrow" if idx % 2 == 0 else "oddrow"
            self.hist_season_tree.insert("", tk.END, iid=str(idx), values=row_vals, tags=(tag,))
        self.hist_season_tree.tag_configure("evenrow", background="#1e293b")
        self.hist_season_tree.tag_configure("oddrow", background="#162032")

        self.show_screen('history')

    def on_select_history_season(self, event):
        sel = self.hist_season_tree.selection()
        if not sel: return
        idx = int(sel[0])
        h = self.active_profile['history'][idx]
        
        hist_rows = []
        for r in h.get('races', []):
            hist_rows.append((r['event'], f"P{r['q']}", f"P{r['f']}", f"P{r['c']}", r.get('grid', 20)))

        for i in self.hist_race_tree.get_children(): self.hist_race_tree.delete(i)
        for r_idx, row_vals in enumerate(hist_rows):
            tag = "evenrow" if r_idx % 2 == 0 else "oddrow"
            self.hist_race_tree.insert("", tk.END, values=row_vals, tags=(tag,))
        self.hist_race_tree.tag_configure("evenrow", background="#1e293b")
        self.hist_race_tree.tag_configure("oddrow", background="#162032")

    # ------------------ PROFILE CREATION ------------------
    def open_new_career_dialog(self):
        if not engine.is_loaded:
            messagebox.showwarning("Load Blueprint First", "Please import and select an active Excel blueprint from the library first!")
            return

        dlg = tk.Toplevel(self)
        dlg.title("Start New Career Profile")
        dlg.geometry("780x540")
        dlg.configure(bg="#090d16")
        
        tk.Label(dlg, text="Career Profile Name:", font=("Helvetica", 9, "bold"), bg="#090d16", fg="#f8fafc").pack(anchor="w", padx=20, pady=(15, 2))
        ent_name = ttk.Entry(dlg, width=50); ent_name.pack(padx=20)

        tk.Label(dlg, text="Driver Name:", font=("Helvetica", 9, "bold"), bg="#090d16", fg="#f8fafc").pack(anchor="w", padx=20, pady=(8, 2))
        ent_driver = ttk.Entry(dlg, width=50); ent_driver.pack(padx=20)

        tk.Label(dlg, text=f"Active Source Blueprint: {engine.active_blueprint_filename}", font=("Helvetica", 8, "italic"), bg="#090d16", fg="#10b981").pack(anchor="w", padx=20, pady=(6, 2))
        tk.Label(dlg, text="Select Starting Championship:", font=("Helvetica", 9, "bold"), bg="#090d16", fg="#f8fafc").pack(anchor="w", padx=20, pady=(8, 4))

        table_frame = tk.Frame(dlg, bg="#1e293b", highlightbackground="#334155", highlightthickness=1)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 10))

        champ_tree = ttk.Treeview(table_frame, columns=("name", "prestige", "categories", "tags"), show="headings", height=8)
        champ_tree.heading("name", text="Championship Series")
        champ_tree.heading("prestige", text="Prestige")
        champ_tree.heading("categories", text="Car Categories")
        champ_tree.heading("tags", text="Tags")
        
        champ_tree.column("name", width=250, stretch=False)
        champ_tree.column("prestige", width=85, stretch=False)
        champ_tree.column("categories", width=180, stretch=False)
        champ_tree.column("tags", width=200, stretch=True)

        h_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=champ_tree.xview)
        v_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=champ_tree.yview)
        champ_tree.configure(xscrollcommand=h_scroll.set, yscrollcommand=v_scroll.set)

        champ_tree.grid(row=0, column=0, sticky="nsew")
        v_scroll.grid(row=0, column=1, sticky="ns")
        h_scroll.grid(row=1, column=0, sticky="ew")

        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        sorted_champs = sorted(engine.championships, key=lambda x: x['prestige'])
        champ_rows = []
        for c in sorted_champs:
            cats_str = ", ".join(c['categories']) if c['categories'] else "General"
            tags_str = ", ".join(c['tags']) if c['tags'] else "Standard"
            champ_rows.append((c['name'], f"Prestige {c['prestige']}", cats_str, tags_str))

        for idx, row_vals in enumerate(champ_rows):
            tag = "evenrow" if idx % 2 == 0 else "oddrow"
            champ_tree.insert("", tk.END, iid=str(idx), values=row_vals, tags=(tag,))
        champ_tree.tag_configure("evenrow", background="#1e293b")
        champ_tree.tag_configure("oddrow", background="#162032")

        if sorted_champs:
            champ_tree.selection_set("0")

        def prompt_car_selection(matching_cars, callback):
            car_dlg = tk.Toplevel(dlg)
            car_dlg.title("Select Team & Car Configuration")
            car_dlg.geometry("540x380")
            car_dlg.configure(bg="#090d16")
            car_dlg.transient(dlg)
            car_dlg.grab_set()

            tk.Label(car_dlg, text="Choose Your Starting Team & Vehicle", font=("Helvetica", 11, "bold"), bg="#090d16", fg="#f8fafc").pack(pady=(15, 5))
            tk.Label(car_dlg, text="Select your preferred team and car combination for this championship:", font=("Helvetica", 9), bg="#090d16", fg="#94a3b8").pack(pady=(0, 10))

            car_table_frame = tk.Frame(car_dlg, bg="#1e293b", highlightbackground="#334155", highlightthickness=1)
            car_table_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 10))

            car_tree = ttk.Treeview(car_table_frame, columns=("class", "team", "model"), show="headings", height=8)
            car_tree.heading("class", text="Car Class")
            car_tree.heading("team", text="Team Name")
            car_tree.heading("model", text="Car Model")

            car_tree.column("class", width=120, stretch=False)
            car_tree.column("team", width=190, stretch=False)
            car_tree.column("model", width=190, stretch=True)

            car_v_scroll = ttk.Scrollbar(car_table_frame, orient="vertical", command=car_tree.yview)
            car_tree.configure(yscrollcommand=car_v_scroll.set)

            car_tree.grid(row=0, column=0, sticky="nsew")
            car_v_scroll.grid(row=0, column=1, sticky="ns")
            car_table_frame.grid_rowconfigure(0, weight=1)
            car_table_frame.grid_columnconfigure(0, weight=1)

            car_rows = []
            for car in matching_cars:
                car_rows.append((car.get('category', 'General'), car['team_name'], car['car_model']))

            for idx, row_vals in enumerate(car_rows):
                tag = "evenrow" if idx % 2 == 0 else "oddrow"
                car_tree.insert("", tk.END, iid=str(idx), values=row_vals, tags=(tag,))
            car_tree.tag_configure("evenrow", background="#1e293b")
            car_tree.tag_configure("oddrow", background="#162032")

            if matching_cars:
                car_tree.selection_set("0")

            def confirm_car_choice():
                sel_car = car_tree.selection()
                if not sel_car:
                    messagebox.showwarning("Selection Required", "Please select a team and car combination.")
                    return
                chosen_car = matching_cars[int(sel_car[0])]
                car_dlg.destroy()
                callback(chosen_car)

            tk.Button(car_dlg, text="Confirm Selection & Start Career", command=confirm_car_choice, bg="#059669", fg="white", font=("Helvetica", 10, "bold"), pady=6, relief=tk.FLAT).pack(fill=tk.X, padx=20, pady=(0, 15))

        def create():
            c_name = ent_name.get()
            d_name = ent_driver.get()
            if not c_name or not d_name: 
                messagebox.showwarning("Missing Fields", "Please enter both a career profile name and driver name.")
                return
            
            sel = champ_tree.selection()
            if not sel:
                messagebox.showwarning("Selection Required", "Please select a starting championship from the table.")
                return

            selected_idx = int(sel[0])
            start_champ = sorted_champs[selected_idx]
            ch_name = start_champ['name']
            
            cal = engine.generate_calendar_for_champ(ch_name)
            
            champ_cats = [cat.lower().strip() for cat in start_champ.get('categories', [])]
            matching_cars = [car for car in engine.cars if any(cc == car['category'].lower().strip() or cc in car['category'].lower().strip() for cc in champ_cats)]
            if not matching_cars:
                matching_cars = [{"team_name": "Factory Squad", "category": "Spec", "car_model": "Spec Race Car"}]

            def finalize_career_creation(chosen_car):
                team_name = chosen_car['team_name']
                car_model = chosen_car['car_model']
                car_class = chosen_car.get('category', 'General')

                c_id = f"prof_{random.randint(10000, 99999)}"
                initial_contract = {
                    "champName": ch_name, 
                    "teamName": team_name, 
                    "carModel": car_model, 
                    "carClass": car_class, 
                    "prestige": start_champ['prestige']
                }
                
                conn = sqlite3.connect(DB_FILE)
                c_db = conn.cursor()
                c_db.execute("INSERT INTO careers VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", 
                             (c_id, c_name, d_name, 1, json.dumps(initial_contract), "[]", json.dumps(cal), "[]", engine.active_blueprint_filename))
                conn.commit()
                conn.close()

                dlg.destroy()
                self.load_profiles_list()

            prompt_car_selection(matching_cars, finalize_career_creation)

        tk.Button(dlg, text="Initialize Career", command=create, bg="#2563eb", fg="white", font=("Helvetica", 10, "bold"), pady=6, relief=tk.FLAT).pack(fill=tk.X, padx=20, pady=(0, 15))

    def load_selected_career(self):
        sel = self.tree.selection()
        if not sel: return
        conn = sqlite3.connect(DB_FILE)
        c = conn.cursor()
        c.execute("SELECT id, name, driver_name, season, current_contract, races, calendar, history, blueprint_name FROM careers WHERE id = ?", (sel[0],))
        r = c.fetchone()
        conn.close()
        if r:
            bp_name = r[8]
            if bp_name:
                bp_path = os.path.join(BLUEPRINTS_DIR, bp_name)
                if os.path.exists(bp_path):
                    engine.load_from_excel(bp_path)
                    self.bp_combo.set(bp_name)
                    self.update_blueprint_width_and_label(bp_name)
                    self.status_lbl.config(text=f"Blueprint Status: Active '{bp_name}'", fg="#10b981")
                else:
                    messagebox.showwarning("Blueprint Missing", f"The blueprint file '{bp_name}' associated with this career was not found in your library folder.\nPlease import or place it back into the AppData blueprints folder.")

            self.active_profile = {
                "id": r[0], "name": r[1], "driverName": r[2], "season": r[3],
                "currentContract": json.loads(r[4]), "races": json.loads(r[5]), "calendar": json.loads(r[6]),
                "history": json.loads(r[7]) if len(r) > 7 and r[7] else [],
                "blueprint_name": bp_name
            }
            self.refresh_dashboard()
            self.show_screen('dashboard')

if __name__ == "__main__":
    app = AMS2CareerApp()
    app.mainloop()